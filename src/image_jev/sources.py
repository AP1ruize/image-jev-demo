from __future__ import annotations

import csv
import hashlib
import io
import json
from pathlib import Path
import zipfile
from typing import Any, Iterator

from .io import sha256_file
from .runtime import load_local_config


def split_name(group_id: str, *, seed: str, train: float, validation: float) -> str:
    fraction = int.from_bytes(hashlib.sha256(f"{seed}\0{group_id}".encode()).digest()[:8], "big") / 2**64
    if fraction < train:
        return "train"
    if fraction < train + validation:
        return "validation"
    return "test"


def rank_key(source_dataset: str, source_id: str, seed: str) -> str:
    return hashlib.sha256(f"{seed}\0{source_dataset}\0{source_id}".encode()).hexdigest()


def _resolve_spaq_layout(root: Path) -> tuple[Path, Path]:
    candidates = [
        (root / "Annotations" / "MOS and Image attribute scores.xlsx", root / "SPAQ zip" / "SPAQ" / "TestImage"),
        (root / "MOS and Image attribute scores.xlsx", root / "TestImage"),
        (
            root / "google-drive" / "Annotations" / "MOS and Image attribute scores.xlsx",
            root / "google-drive" / "SPAQ zip" / "SPAQ" / "TestImage",
        ),
        (root / "annotations" / "MOS and Image attribute scores.xlsx", root / "images"),
    ]
    for scores, images in candidates:
        if scores.is_file() and images.is_dir():
            return scores, images
    raise FileNotFoundError(
        f"could not find SPAQ annotations + TestImage under {root}; see docs/DATASETS.md"
    )


def _load_xlsx_rows(path: Path) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = workbook.worksheets[0].iter_rows(values_only=True)
        headers = [str(value) for value in next(rows)]
        return [dict(zip(headers, row, strict=True)) for row in rows]
    finally:
        workbook.close()


def ingest_spaq_from_root(root: Path) -> Iterator[dict[str, Any]]:
    scores_path, image_dir = _resolve_spaq_layout(root)
    by_name = {str(row["Image name"]): row for row in _load_xlsx_rows(scores_path)}
    for path in sorted(image_dir.glob("*.jpg")):
        row = by_name.get(path.name)
        if row is None:
            continue
        yield {
            "source_dataset": "spaq",
            "source_id": path.name,
            "group_id": f"spaq:{path.name}",
            "path": str(path),
            "labels": {
                "MOS": float(row["MOS"]),
                "Brightness": float(row["Brightness"]),
                "Colorfulness": float(row["Colorfulness"]),
                "Contrast": float(row["Contrast"]),
                "Noisiness": float(row["Noisiness"]),
                "Sharpness": float(row["Sharpness"]),
            },
        }


def ingest_spaq_from_manifest(manifest: Path, path_root: Path | None) -> Iterator[dict[str, Any]]:
    base = path_root or manifest.parent
    for line in manifest.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        row = json.loads(line)
        asset = row.get("asset") or {}
        rel = asset.get("path") or asset.get("image_name")
        if not rel:
            continue
        path = Path(rel)
        if not path.is_absolute():
            path = base / path
        scores = row.get("scores") or row.get("labels") or {}
        if "MOS" not in scores:
            continue
        yield {
            "source_dataset": "spaq",
            "source_id": row.get("source_id") or path.name,
            "group_id": row.get("group_id") or f"spaq:{path.name}",
            "path": str(path),
            "sha256": asset.get("sha256"),
            "labels": {
                "MOS": float(scores["MOS"]),
                "Brightness": float(scores.get("Brightness", 0.0)),
                "Colorfulness": float(scores.get("Colorfulness", 0.0)),
                "Contrast": float(scores.get("Contrast", 0.0)),
                "Noisiness": float(scores.get("Noisiness", 0.0)),
                "Sharpness": float(scores.get("Sharpness", 0.0)),
            },
        }


def ingest_koniq(images_zip: Path, scores_zip: Path) -> Iterator[dict[str, Any]]:
    with zipfile.ZipFile(images_zip) as archive:
        members = {Path(name).name: name for name in archive.namelist() if name.lower().endswith((".jpg", ".jpeg"))}
    with zipfile.ZipFile(scores_zip) as archive:
        csv_name = next(name for name in archive.namelist() if name.lower().endswith(".csv"))
        text = archive.read(csv_name).decode("utf-8-sig")
    for row in csv.DictReader(io.StringIO(text)):
        source_id = row["image_name"]
        member = members.get(source_id)
        if member is None:
            continue
        yield {
            "source_dataset": "koniq10k",
            "source_id": source_id,
            "group_id": f"koniq10k:{source_id}",
            "archive_path": str(images_zip),
            "archive_member": member,
            "labels": {
                "MOS": float(row["MOS"]),
                "MOS_zscore": float(row.get("MOS_zscore") or 0.0),
                "SD": float(row.get("SD") or 0.0),
            },
        }


def collect_sources(local: dict[str, Any] | None = None) -> list[dict[str, Any]]:
    cfg = local if local is not None else load_local_config()
    sources = cfg.get("sources") or {}
    records: list[dict[str, Any]] = []
    manifest = sources.get("spaq_manifest")
    root = sources.get("spaq_root")
    if manifest:
        path_root = Path(sources["spaq_path_root"]) if sources.get("spaq_path_root") else None
        records.extend(ingest_spaq_from_manifest(Path(manifest), path_root))
    elif root:
        records.extend(ingest_spaq_from_root(Path(root)))
    images_zip = sources.get("koniq_images_zip")
    scores_zip = sources.get("koniq_scores_zip")
    if images_zip and scores_zip:
        records.extend(ingest_koniq(Path(images_zip), Path(scores_zip)))
    return records


def copy_or_extract(record: dict[str, Any], dest: Path) -> dict[str, str]:
    dest.parent.mkdir(parents=True, exist_ok=True)
    if record.get("archive_path"):
        with zipfile.ZipFile(record["archive_path"]) as archive:
            with archive.open(record["archive_member"]) as source, dest.open("wb") as output:
                output.write(source.read())
    else:
        src = Path(record["path"])
        if not src.is_file():
            raise FileNotFoundError(src)
        dest.write_bytes(src.read_bytes())
    digest = record.get("sha256") or sha256_file(dest)
    if sha256_file(dest) != digest and record.get("sha256"):
        raise ValueError(f"sha256 mismatch after copy: {dest}")
    return {"asset_id": f"{record['source_dataset']}:{record['source_id']}", "path": dest.as_posix(), "sha256": sha256_file(dest)}
