#!/usr/bin/env python3
"""Build datasets/inventory.json and catalog.jsonl. Does not assign train/test splits."""
from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DS = ROOT / "datasets"


def _json_dump(path: Path, obj: Any) -> None:
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def _rel(path: Path) -> str:
    return path.relative_to(ROOT).as_posix()


def spaq_rows() -> list[dict[str, Any]]:
    scores = DS / "spaq" / "annotations" / "MOS and Image attribute scores.xlsx"
    images = DS / "spaq" / "images"
    wb = load_workbook(scores, read_only=True, data_only=True)
    try:
        rows = wb.worksheets[0].iter_rows(values_only=True)
        headers = [str(v) for v in next(rows)]
        table = [dict(zip(headers, row, strict=True)) for row in rows]
    finally:
        wb.close()
    by_name = {str(r["Image name"]): r for r in table}
    out = []
    for path in sorted(images.glob("*.jpg")):
        row = by_name.get(path.name)
        if row is None:
            continue
        out.append(
            {
                "source_dataset": "spaq",
                "source_id": path.name,
                "path": _rel(path),
                "labels": {
                    "MOS": float(row["MOS"]),
                    "Brightness": float(row["Brightness"]),
                    "Colorfulness": float(row["Colorfulness"]),
                    "Contrast": float(row["Contrast"]),
                    "Noisiness": float(row["Noisiness"]),
                    "Sharpness": float(row["Sharpness"]),
                },
            }
        )
    return out


def koniq_rows() -> tuple[list[dict[str, Any]], dict[str, int]]:
    scores = DS / "koniq10k" / "labels" / "koniq10k_scores_and_distributions.csv"
    plus = DS / "koniq10k" / "labels" / "koniq++database_distributions.csv"
    images = DS / "koniq10k" / "images"
    score_rows = {r["image_name"]: r for r in csv.DictReader(scores.open(encoding="utf-8-sig"))}
    plus_rows = {r["filename"]: r for r in csv.DictReader(plus.open(encoding="utf-8-sig"))}
    jpg = {p.name: p for p in images.glob("*.jpg")}
    out = []
    for name, row in sorted(score_rows.items()):
        path = jpg.get(name)
        if path is None:
            continue
        item: dict[str, Any] = {
            "source_dataset": "koniq10k",
            "source_id": name,
            "path": _rel(path),
            "labels": {
                "MOS": float(row["MOS"]),
                "MOS_zscore": float(row.get("MOS_zscore") or 0.0),
                "SD": float(row.get("SD") or 0.0),
            },
        }
        extra = plus_rows.get(name)
        if extra:
            item["labels"]["koniq++"] = {
                "qmos": float(extra["qmos"]),
                "artifacts": float(extra["artifacts"]),
                "blur": float(extra["blur"]),
                "contrast": float(extra["contrast"]),
                "colors": float(extra["colors"]),
                "other": float(extra["other"]),
                "degraded.amount": float(extra["degraded.amount"]),
            }
        out.append(item)
    stats = {
        "zip_jpegs": len(jpg),
        "score_csv_ids": len(score_rows),
        "joined": len(out),
        "koniq++_joined": sum(1 for i in out if "koniq++" in i["labels"]),
        "unscored_extra_jpegs": len(jpg) - len(out),
    }
    return out, stats


def cid2013_rows() -> list[dict[str, Any]]:
    xlsx = DS / "cid2013" / "labels" / "CID2013 data - version 12112014.xlsx"
    images = DS / "cid2013" / "images"
    by_stem = {p.stem: p for p in images.rglob("*.jpg")}
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    try:
        ws = wb["CID2013 MOS"]
        rows = ws.iter_rows(values_only=True)
        headers = [str(v) if v is not None else f"col_{i}" for i, v in enumerate(next(rows))]
        table = [dict(zip(headers, row, strict=False)) for row in rows]
    finally:
        wb.close()
    out = []
    for row in table:
        source_id = str(row["Source_ID"]).strip()
        path = by_stem.get(source_id)
        if path is None:
            continue
        out.append(
            {
                "source_dataset": "cid2013",
                "source_id": source_id,
                "path": _rel(path),
                "labels": {
                    "MOS_set": float(row["Image set specific  MOS"]),
                    "MOS_realigned": float(row["Realigned MOS"]),
                    "Sharpness_MOS": float(row["Sharpness MOS"] or 0.0),
                    "Graininess_MOS": float(row["Graininess MOS"] or 0.0),
                },
            }
        )
    return out


def uhd_rows() -> list[dict[str, Any]]:
    root = DS / "uhd-iqa"
    meta = root / "uhd-iqa-metadata.csv"
    if not meta.is_file():
        nested = root / "UHD-IQA-database" / "uhd-iqa-metadata.csv"
        meta = nested if nested.is_file() else meta
    if not meta.is_file():
        return []
    by_name: dict[str, Path] = {}
    search_root = meta.parent
    for path in search_root.rglob("*.jpg"):
        by_name[path.name] = path
    out = []
    with meta.open(encoding="utf-8-sig", newline="") as fh:
        for row in csv.DictReader(fh):
            name = row.get("image_name") or row.get("filename")
            if not name:
                continue
            if not name.lower().endswith(".jpg"):
                name = f"{name}.jpg"
            path = by_name.get(name) or by_name.get(Path(name).name)
            if path is None:
                continue
            labels: dict[str, Any] = {}
            if row.get("quality_mos"):
                labels["MOS"] = float(row["quality_mos"])
            if row.get("set"):
                labels["source_set"] = row["set"]
            out.append(
                {
                    "source_dataset": "uhd-iqa",
                    "source_id": Path(name).name,
                    "path": _rel(path),
                    "labels": labels,
                }
            )
    return out


def main() -> None:
    catalog: list[dict[str, Any]] = []
    inventory: dict[str, Any] = {"split": None, "sources": {}}

    spaq = spaq_rows()
    catalog.extend(spaq)
    inventory["sources"]["spaq"] = {
        "images": len(spaq),
        "layout": "datasets/spaq/images is an NTFS junction to AlphaHalf SPAQ TestImage",
        "annotations": "datasets/spaq/annotations",
    }

    koniq, kstats = koniq_rows()
    catalog.extend(koniq)
    inventory["sources"]["koniq10k"] = kstats
    inventory["sources"]["koniq++"] = {
        "label_file": "datasets/koniq10k/labels/koniq++database_distributions.csv",
        "joined_to_koniq10k": kstats["koniq++_joined"],
    }

    cid = cid2013_rows()
    catalog.extend(cid)
    inventory["sources"]["cid2013"] = {"images_joined_to_MOS_sheet": len(cid)}

    uhd = uhd_rows()
    if uhd:
        catalog.extend(uhd)
        inventory["sources"]["uhd-iqa"] = {"images": len(uhd)}
    else:
        inventory["sources"]["uhd-iqa"] = {"status": "not_on_d_drive"}

    inventory["n"] = len(catalog)
    _json_dump(DS / "inventory.json", inventory)
    with (DS / "catalog.jsonl").open("w", encoding="utf-8") as fh:
        for row in catalog:
            fh.write(json.dumps(row, ensure_ascii=False) + "\n")
    print(json.dumps(inventory, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
